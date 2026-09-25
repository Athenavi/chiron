# RAG 构建器 — 支持双向量数据库 + 混合解析器 + LangChain 分块
from __future__ import annotations

import asyncio
import logging
import re
import uuid
from collections.abc import AsyncIterator
from enum import StrEnum
from typing import Any

from app.config import settings
from app.gateway.router import GatewayRouter
from app.interfaces.vectorstore import VectorStore

logger = logging.getLogger(__name__)


class VectorDBType(StrEnum):
    """向量数据库类型"""

    MILVUS = "milvus"
    PGVECTOR = "pgvector"


class ParserType(StrEnum):
    """解析器类型"""

    UNSTRUCTURED = "unstructured"
    MARKITDOWN = "markitdown"
    LIGHTWEIGHT = "lightweight"


class RAGBuilder:
    """RAG 知识库构建器 — 支持多种配置"""

    def __init__(
        self,
        llm_gateway: GatewayRouter,
        vector_store: VectorStore | None = None,
        vector_db_type: VectorDBType | str | None = None,
        parser_type: ParserType = ParserType.UNSTRUCTURED,
    ):
        self._gateway = llm_gateway
        self._vector_store = vector_store
        self._parser_type = parser_type
        self._milvus_connected = False
        self._local_encoder: Any = None
        self._local_encoder_lock = asyncio.Lock()
        self._pg_pool = None
        # 体内求值：非法配置回退默认值，避免模块 import 崩溃
        if vector_db_type is None:
            vector_db_type = settings.vector_db_type
        try:
            self._vector_db_type = (
                vector_db_type
                if isinstance(vector_db_type, VectorDBType)
                else VectorDBType(vector_db_type)
            )
        except ValueError:
            logger.warning("未知 vector_db_type=%r，回退 milvus", vector_db_type)
            self._vector_db_type = VectorDBType.MILVUS

    def _ensure_milvus(self) -> None:
        """确保 Milvus 连接"""
        if not self._milvus_connected:
            from pymilvus import connections

            host = settings.milvus_address.split(":")[0]
            port = (
                int(settings.milvus_address.split(":")[1])
                if ":" in settings.milvus_address
                else 19530
            )
            connections.connect(alias="default", host=host, port=port)
            self._milvus_connected = True
            logger.info("Milvus connected: %s:%d", host, port)

    async def build_document(
        self,
        kb_id: str,
        doc_id: str,
        content: bytes,
        file_type: str,
        filename: str,
        tenant_id: str,
        vector_db: str | None = None,
        parser_type: str | None = None,
    ) -> AsyncIterator[dict[str, Any]]:
        """
        构建单个文档的 RAG 索引（SSE 流式返回进度）
        """
        # 使用指定的向量数据库类型，或默认；非法值回退默认，避免 500
        if vector_db:
            try:
                db_type = VectorDBType(vector_db)
            except ValueError:
                logger.warning(
                    "未知 vector_db=%r，使用默认 %s",
                    vector_db,
                    self._vector_db_type.value,
                )
                db_type = self._vector_db_type
        else:
            db_type = self._vector_db_type
        parse_type = ParserType(parser_type) if parser_type else self._parser_type

        try:
            # Step 1: 解析
            yield {"type": "progress", "step": "parsing", "progress": 0.1}
            parse_result = self._parse_document(
                content, file_type, filename, parse_type
            )
            if parse_result.get("error"):
                yield {"type": "error", "message": parse_result["error"]}
                return

            text = parse_result["text"]
            char_count = parse_result["char_count"]
            if not text.strip():
                yield {"type": "error", "message": "文档内容为空"}
                return

            # Step 2: 分块
            yield {"type": "progress", "step": "chunking", "progress": 0.3}
            chunks = self._chunk_text(text)
            if not chunks:
                yield {"type": "error", "message": "分块结果为空"}
                return

            # Step 3: 嵌入
            yield {"type": "progress", "step": "embedding", "progress": 0.5}
            embeddings = await self._compute_embeddings(chunks)
            yield {"type": "progress", "step": "embedding", "progress": 0.8}

            # Step 4: 存储
            yield {"type": "progress", "step": "storing", "progress": 0.9}
            await self._store_vectors(
                kb_id, doc_id, tenant_id, chunks, embeddings, db_type
            )

            yield {"type": "progress", "step": "storing", "progress": 1.0}
            yield {
                "type": "complete",
                "chunk_count": len(chunks),
                "char_count": char_count,
                "page_count": parse_result.get("page_count", 0),
            }
        except Exception as e:
            logger.error("文档构建失败: %s", e)
            yield {"type": "error", "message": str(e)}

    def _parse_document(
        self,
        content: bytes,
        file_type: str,
        filename: str,
        parser_type: ParserType,
    ) -> dict[str, Any]:
        """解析文档"""
        if parser_type == ParserType.UNSTRUCTURED:
            return self._parse_with_unstructured(content, file_type, filename)
        elif parser_type == ParserType.MARKITDOWN:
            return self._parse_with_markitdown(content, file_type, filename)
        elif parser_type == ParserType.LIGHTWEIGHT:
            return self._parse_lightweight(content, file_type, filename)
        else:
            return {"error": f"不支持的解析器类型: {parser_type}"}

    def _parse_with_unstructured(
        self,
        content: bytes,
        file_type: str,
        filename: str,
    ) -> dict[str, Any]:
        """使用 Unstructured 解析"""
        try:
            import os

            # 使用通用 partition 函数自动处理所有文件类型
            import tempfile

            from unstructured.partition.auto import partition

            suffix = f".{file_type}" if file_type else ""
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            try:
                elements = partition(filename=tmp_path)
            finally:
                os.unlink(tmp_path)

            text = "\n\n".join([str(el) for el in elements])
            return {
                "text": text,
                "char_count": len(text),
                "page_count": len(
                    [el for el in elements if type(el).__name__ == "PageBreak"]
                ),
            }
        except Exception as e:
            logger.warning("Unstructured 解析失败: %s", e)
            return {"error": str(e)}

    def _parse_with_markitdown(
        self,
        content: bytes,
        file_type: str,
        filename: str,
    ) -> dict[str, Any]:
        """使用 markitdown 解析"""
        try:
            from markitdown import MarkItDown

            md = MarkItDown()
            result = md.convert(content)
            return {
                "text": result.text_content,
                "char_count": len(result.text_content),
                "page_count": 1,
            }
        except Exception as e:
            logger.warning("markitdown 解析失败: %s", e)
            return {"error": str(e)}

    def _parse_lightweight(
        self,
        content: bytes,
        file_type: str,
        filename: str,
    ) -> dict[str, Any]:
        """轻量级解析"""
        try:
            if file_type == "pdf":
                return self._parse_pdf_lightweight(content)
            elif file_type in ("docx", "doc"):
                return self._parse_docx_lightweight(content)
            elif file_type in ("xlsx", "xls"):
                return self._parse_xlsx_lightweight(content)
            elif file_type == "html":
                return self._parse_html_lightweight(content)
            elif file_type == "markdown":
                return {
                    "text": content.decode("utf-8"),
                    "char_count": len(content),
                    "page_count": 1,
                }
            else:
                return {"error": f"轻量级解析器不支持: {file_type}"}
        except Exception as e:
            logger.warning("轻量级解析失败: %s", e)
            return {"error": str(e)}

    def _parse_pdf_lightweight(self, content: bytes) -> dict[str, Any]:
        """使用 PyMuPDF 解析 PDF"""
        import fitz

        doc = fitz.open(stream=content, filetype="pdf")
        try:
            text_parts = []
            for page in doc:
                text_parts.append(page.get_text())
            text = "\n\n".join(text_parts)
            return {
                "text": text,
                "char_count": len(text),
                "page_count": len(doc),
            }
        finally:
            doc.close()

    def _parse_docx_lightweight(self, content: bytes) -> dict[str, Any]:
        """使用 python-docx 解析 Word"""
        import io

        from docx import Document

        doc: Any = Document(io.BytesIO(content))
        try:
            text = "\n\n".join(
                [para.text for para in doc.paragraphs if para.text.strip()]
            )
            return {
                "text": text,
                "char_count": len(text),
                "page_count": 1,
            }
        finally:
            doc.close()

    def _parse_xlsx_lightweight(self, content: bytes) -> dict[str, Any]:
        """使用 openpyxl 解析 Excel"""
        import io

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content))
        try:
            text_parts = []
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(
                        [str(cell) for cell in row if cell is not None]
                    )
                    if row_text.strip():
                        text_parts.append(row_text)
            text = "\n".join(text_parts)
            return {
                "text": text,
                "char_count": len(text),
                "page_count": len(wb.sheetnames),
            }
        finally:
            wb.close()

    def _parse_html_lightweight(self, content: bytes) -> dict[str, Any]:
        """使用 BeautifulSoup 解析 HTML"""
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(content, "html.parser")
        text = soup.get_text(separator="\n", strip=True)
        return {
            "text": text,
            "char_count": len(text),
            "page_count": 1,
        }

    def _chunk_text(self, text: str) -> list[dict[str, Any]]:
        """使用 LangChain Text Splitter 分块"""
        try:
            from langchain.text_splitter import RecursiveCharacterTextSplitter

            splitter = RecursiveCharacterTextSplitter(
                chunk_size=settings.chunk_size,
                chunk_overlap=settings.chunk_overlap,
                length_function=len,
                separators=["\n\n", "\n", "。", ".", "!", "！", "?", "？", " "],
            )
            chunks = splitter.split_text(text)
            return [{"index": i, "content": chunk} for i, chunk in enumerate(chunks)]
        except Exception as e:
            logger.warning("LangChain 分块失败: %s", e)
            # fallback 到简单分块
            return self._chunk_text_simple(text)

    def _chunk_text_simple(self, text: str) -> list[dict[str, Any]]:
        """简单分块（fallback）"""
        chunks = []
        chunk_size = settings.chunk_size
        chunk_overlap = settings.chunk_overlap

        start = 0
        index = 0
        while start < len(text):
            end = start + chunk_size
            chunk = text[start:end]
            if chunk.strip():
                chunks.append({"index": index, "content": chunk})
                index += 1
            start = end - chunk_overlap

        return chunks

    async def _embed_text(self, text: str) -> list[float] | None:
        """统一嵌入入口：本地优先、API 回退（存储与查询口径一致）"""
        embedding = await self._get_local_embedding(text)
        if embedding is None:
            embedding = await self._get_api_embedding(text)
        return embedding

    async def _compute_embeddings(self, chunks: list[dict[str, Any]]) -> list[list[float] | None]:
        """计算嵌入向量（本地优先，API 回退）"""
        embeddings = []
        for i, chunk in enumerate(chunks):
            try:
                embeddings.append(await self._embed_text(chunk["content"]))
            except Exception as e:
                logger.warning("嵌入计算失败 chunk %d: %s", i, e)
                embeddings.append(None)
        return embeddings

    async def _get_local_embedding(self, text: str) -> list[float] | None:
        """使用本地模型计算嵌入（可插拔，默认关闭）

        配置 settings.local_embedding_model（如 BGE/Jina 的本地路径/模型名）后启用；
        依赖 sentence-transformers（可选 extra）。任何失败均回退到 API 嵌入。
        """
        if not settings.local_embedding_model:
            return None
        try:
            async with self._local_encoder_lock:
                if self._local_encoder is None:
                    from sentence_transformers import SentenceTransformer

                    # 经 Any 变量转交：sentence_transformers 的 stub 不完整，直接传类会让
                    # mypy 认为 to_thread 的返回是 None（它按 __init__ 的返回推断）。
                    encoder_cls: Any = SentenceTransformer
                    self._local_encoder = await asyncio.to_thread(
                        encoder_cls, settings.local_embedding_model
                    )
            vector = await asyncio.to_thread(self._local_encoder.encode, text)
            return [float(x) for x in vector.tolist()]
        except Exception as e:
            logger.warning("本地嵌入计算失败（回退到 API）: %s", e)
            return None

    async def _get_api_embedding(self, text: str) -> list[float] | None:
        """使用 API 计算嵌入"""
        try:
            resp = await self._gateway.embed(text, settings.embedding_model)
            return resp.embedding
        except Exception as e:
            logger.warning("API 嵌入计算失败: %s", e)
            return None

    async def _store_vectors(
        self,
        kb_id: str,
        doc_id: str,
        tenant_id: str,
        chunks: list[dict[str, Any]],
        embeddings: list[list[float] | None],
        db_type: VectorDBType,
    ) -> None:
        """存储向量"""
        if db_type == VectorDBType.MILVUS:
            await self._store_milvus(kb_id, doc_id, tenant_id, chunks, embeddings)
        elif db_type == VectorDBType.PGVECTOR:
            await self._store_pgvector(kb_id, doc_id, tenant_id, chunks, embeddings)
        else:
            raise ValueError(f"不支持的向量数据库: {db_type}")

    async def _store_milvus(
        self,
        kb_id: str,
        doc_id: str,
        tenant_id: str,
        chunks: list[dict[str, Any]],
        embeddings: list[list[float] | None],
    ) -> None:
        """存储到 Milvus"""
        if self._vector_store:
            # 使用 VectorStore 接口
            await self._store_with_interface(
                kb_id, doc_id, tenant_id, chunks, embeddings
            )
        else:
            # 直接使用 pymilvus
            await self._store_milvus_direct(
                kb_id, doc_id, tenant_id, chunks, embeddings
            )

    async def _store_with_interface(
        self,
        kb_id: str,
        doc_id: str,
        tenant_id: str,
        chunks: list[dict[str, Any]],
        embeddings: list[list[float] | None],
    ) -> None:
        """使用 VectorStore 接口存储"""
        store = self._vector_store
        if store is None:  # 调用方只在 _vector_store 非空时进来；这里收窄类型
            raise ValueError("VectorStore 未初始化")
        collection = f"kb_{kb_id.replace('-', '_')}"
        dim = (
            len(embeddings[0])
            if embeddings and embeddings[0]
            else settings.embedding_dim
        )

        await store.ensure_collection(collection, dim)

        ids = []
        vectors = []
        payloads = []
        for _i, (chunk, embedding) in enumerate(zip(chunks, embeddings, strict=False)):
            if embedding is None:
                continue
            ids.append(str(uuid.uuid4()))
            vectors.append(embedding)
            payloads.append(
                {
                    "content": chunk["content"][:65000],
                    "kb_id": kb_id,
                    "doc_id": doc_id,
                    "tenant_id": tenant_id,
                    "chunk_index": chunk["index"],
                }
            )

        if ids:
            await store.insert(collection, ids, vectors, payloads)
        else:
            raise ValueError("无有效向量可存储，VectorStore 存储失败")
        logger.info("VectorStore 存储完成: %d chunks", len(ids))

    async def _store_milvus_direct(
        self,
        kb_id: str,
        doc_id: str,
        tenant_id: str,
        chunks: list[dict[str, Any]],
        embeddings: list[list[float] | None],
    ) -> None:
        """直接存储到 Milvus"""
        from pymilvus import Collection, CollectionSchema, DataType, FieldSchema

        self._ensure_milvus()
        dim = (
            len(embeddings[0])
            if embeddings and embeddings[0]
            else settings.embedding_dim
        )

        fields = [
            FieldSchema(
                name="id", dtype=DataType.VARCHAR, is_primary=True, max_length=64
            ),
            FieldSchema(name="kb_id", dtype=DataType.VARCHAR, max_length=64),
            FieldSchema(name="doc_id", dtype=DataType.VARCHAR, max_length=64),
            FieldSchema(name="tenant_id", dtype=DataType.VARCHAR, max_length=64),
            FieldSchema(name="chunk_index", dtype=DataType.INT64),
            FieldSchema(name="content", dtype=DataType.VARCHAR, max_length=65535),
            FieldSchema(name="embedding", dtype=DataType.FLOAT_VECTOR, dim=dim),
        ]
        schema = CollectionSchema(fields, description="RAG knowledge chunks")
        collection_name = f"kb_{kb_id.replace('-', '_')}"
        try:
            collection = Collection(collection_name)
        except Exception:
            collection = Collection(collection_name, schema)
            collection.create_index(
                "embedding",
                {
                    "metric_type": "COSINE",
                    "index_type": "IVF_FLAT",
                    "params": {"nlist": 1024},
                },
            )

        ids, kb_ids, doc_ids, tenant_ids, chunk_indices, contents, valid_embeddings = (
            [],
            [],
            [],
            [],
            [],
            [],
            [],
        )
        for _i, (chunk, embedding) in enumerate(zip(chunks, embeddings, strict=False)):
            if embedding is None:
                continue
            ids.append(str(uuid.uuid4()))
            kb_ids.append(kb_id)
            doc_ids.append(doc_id)
            tenant_ids.append(tenant_id)
            chunk_indices.append(chunk["index"])
            contents.append(chunk["content"][:65000])
            valid_embeddings.append(embedding)

        if ids:
            # 重建幂等：先按文档清理旧向量再插入（与 pgvector 路径一致，避免重试累积重复）
            # doc_id 为服务端生成的 uuid，转义引号防御表达式注入
            safe_doc_id = doc_id.replace('"', '\\"')
            await asyncio.to_thread(collection.delete, f'doc_id == "{safe_doc_id}"')
            await asyncio.to_thread(
                collection.insert,
                [
                    ids,
                    kb_ids,
                    doc_ids,
                    tenant_ids,
                    chunk_indices,
                    contents,
                    valid_embeddings,
                ],
            )
            await asyncio.to_thread(collection.flush)
        else:
            raise ValueError("无有效向量可存储，Milvus 存储失败")
        logger.info("Milvus 存储完成: %d chunks", len(ids))

    @staticmethod
    def _validate_table_name(table: str) -> str:
        """校验表名为合法标识符，防止配置注入 SQL"""
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", table):
            raise ValueError(f"非法表名: {table!r}")
        return table

    async def _pg_conn(self) -> tuple[Any, Any]:
        """获取 PostgreSQL 连接与释放函数：统一通过 DBManager"""
        from app.db import get_pool

        pool = get_pool()
        conn = await pool.acquire()
        return conn, pool.release

    async def _store_pgvector(
        self,
        kb_id: str,
        doc_id: str,
        tenant_id: str,
        chunks: list[dict[str, Any]],
        embeddings: list[list[float] | None],
    ) -> None:
        """存储到 PostgreSQL (pgvector)"""
        table = self._validate_table_name(settings.pgvector_table)
        rows = []
        skipped = 0
        for chunk, embedding in zip(chunks, embeddings, strict=False):
            if embedding is None:
                continue
            if len(embedding) != settings.embedding_dim:
                skipped += 1
                logger.warning(
                    "chunk %d 维度 %d 与 embedding_dim=%d 不一致，跳过",
                    chunk["index"],
                    len(embedding),
                    settings.embedding_dim,
                )
                continue
            rows.append(
                (
                    str(uuid.uuid4()),
                    kb_id,
                    doc_id,
                    tenant_id,
                    chunk["index"],
                    chunk["content"][:65000],
                    f"[{','.join(format(float(x), '.8f') for x in embedding)}]",
                )
            )

        if not rows:
            if skipped:
                raise ValueError(
                    f"全部 {skipped} 个向量维度与 embedding_dim={settings.embedding_dim} 不一致，pgvector 存储失败"
                )
            raise ValueError("无有效向量可存储，pgvector 存储失败")

        conn, release = await self._pg_conn()
        try:
            async with conn.transaction():
                # 重建幂等：先按文档清理旧向量再插入，避免文档重建累积重复
                await conn.execute(
                    f"DELETE FROM {table} WHERE document_id = $1", doc_id
                )
                await conn.execute(
                    f"""
                    INSERT INTO {table}
                        (id, knowledge_base_id, document_id, tenant_id, chunk_index, content, embedding)
                    SELECT id, knowledge_base_id, document_id, tenant_id, chunk_index, content, embedding::vector
                    FROM unnest(
                        $1::text[], $2::text[], $3::text[], $4::text[], $5::int[],
                        $6::text[], $7::text[]
                    ) AS t(id, knowledge_base_id, document_id, tenant_id, chunk_index, content, embedding)
                    ON CONFLICT (id) DO NOTHING
                    """,
                    [r[0] for r in rows],
                    [r[1] for r in rows],
                    [r[2] for r in rows],
                    [r[3] for r in rows],
                    [r[4] for r in rows],
                    [r[5] for r in rows],
                    [r[6] for r in rows],
                )
        finally:
            await release(conn)
        logger.info("pgvector 存储完成: %d chunks (doc_id=%s)", len(rows), doc_id)

    async def query(
        self,
        kb_id: str,
        query: str,
        top_k: int = 5,
        threshold: float = 0.5,
        vector_db: str | None = None,
    ) -> list[dict[str, Any]]:
        """查询知识库"""
        if vector_db:
            try:
                db_type = VectorDBType(vector_db)
            except ValueError:
                logger.warning(
                    "未知 vector_db=%r，使用默认 %s",
                    vector_db,
                    self._vector_db_type.value,
                )
                db_type = self._vector_db_type
        else:
            db_type = self._vector_db_type

        if db_type == VectorDBType.MILVUS:
            return await self.query_milvus(kb_id, query, top_k, threshold)
        elif db_type == VectorDBType.PGVECTOR:
            return await self.query_pgvector(kb_id, query, top_k, threshold)
        else:
            raise ValueError(f"不支持的向量数据库: {db_type}")

    async def query_milvus(
        self, kb_id: str, query: str, top_k: int = 5, threshold: float = 0.5
    ) -> list[dict[str, Any]]:
        """从 Milvus 查询"""
        if self._vector_store:
            return await self._query_with_interface(kb_id, query, top_k, threshold)

        from pymilvus import Collection

        self._ensure_milvus()
        collection_name = f"kb_{kb_id.replace('-', '_')}"
        try:
            collection = Collection(collection_name)
            collection.load()
        except Exception as e:
            logger.error("Milvus collection load failed: %s", e)
            return []

        embedding = await self._embed_text(query)
        if not embedding:
            return []

        results = await asyncio.to_thread(
            collection.search,
            data=[embedding],
            anns_field="embedding",
            param={"metric_type": "COSINE", "params": {"nprobe": 10}},
            limit=top_k,
            output_fields=["doc_id", "chunk_index", "content"],
        )
        return [
            {
                "id": hit.id,
                "content": hit.entity.get("content", ""),
                "doc_id": hit.entity.get("doc_id", ""),
                "chunk_index": hit.entity.get("chunk_index", 0),
                "score": hit.score,
            }
            for hits in results
            for hit in hits
            if hit.score >= threshold
        ]

    async def _query_with_interface(
        self, kb_id: str, query: str, top_k: int, threshold: float
    ) -> list[dict[str, Any]]:
        """使用 VectorStore 接口查询"""
        collection = f"kb_{kb_id.replace('-', '_')}"

        embedding = await self._embed_text(query)
        if not embedding:
            return []

        store = self._vector_store
        if store is None:  # 调用方只在 _vector_store 非空时进来；这里收窄类型
            raise ValueError("VectorStore 未初始化")
        results = await store.search(
            collection=collection,
            query_vector=embedding,
            top_k=top_k,
            threshold=threshold,
        )

        return [
            {
                "id": r["id"],
                "content": r.get("content", ""),
                "doc_id": r.get("doc_id", ""),
                "chunk_index": r.get("chunk_index", 0),
                "score": r.get("score", 0),
            }
            for r in results
        ]

    async def query_pgvector(
        self, kb_id: str, query: str, top_k: int = 5, threshold: float = 0.5
    ) -> list[dict[str, Any]]:
        """从 PostgreSQL (pgvector) 查询（余弦相似度，与 Milvus 语义对齐）"""
        query_vector = await self._embed_text(query)
        if not query_vector:
            return []

        table = self._validate_table_name(settings.pgvector_table)
        qv = f"[{','.join(format(float(x), '.8f') for x in query_vector)}]"
        try:
            conn, release = await self._pg_conn()
        except Exception as e:
            logger.error("pgvector 连接失败: %s", e)
            return []
        try:
            rows = await conn.fetch(
                f"""
                SELECT id, document_id, chunk_index, content,
                       1 - (embedding <=> $2::vector) AS score
                FROM {table}
                WHERE knowledge_base_id = $1
                  AND 1 - (embedding <=> $2::vector) >= $3
                ORDER BY embedding <=> $2::vector
                LIMIT $4
                """,
                kb_id,
                qv,
                threshold,
                top_k,
            )
        except Exception as e:
            logger.error("pgvector 查询失败: %s", e)
            return []
        finally:
            await release(conn)

        return [
            {
                "id": row["id"],
                "content": row["content"],
                "doc_id": row["document_id"],
                "chunk_index": row["chunk_index"],
                "score": row["score"],
            }
            for row in rows
        ]

    async def query_qdrant(
        self, kb_id: str, query: str, top_k: int = 5, threshold: float = 0.5
    ) -> list[dict[str, Any]]:
        """从 Qdrant 查询（兼容旧代码）"""
        from qdrant_client import QdrantClient

        client = QdrantClient(host="localhost", port=6333)
        collection_name = f"kb_{kb_id.replace('-', '_')}"
        embedding = await self._embed_text(query)
        if not embedding:
            return []
        results = await asyncio.to_thread(
            client.search,
            collection_name=collection_name,
            query_vector=embedding,
            limit=top_k,
            score_threshold=threshold,
        )
        return [
            {
                "id": hit.id,
                "content": hit.payload.get("content", ""),
                "doc_id": hit.payload.get("doc_id", ""),
                "chunk_index": hit.payload.get("chunk_index", 0),
                "score": hit.score,
            }
            for hit in results
        ]
