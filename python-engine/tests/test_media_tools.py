"""`media_create` 的回归测试：agent 产出文件 → **真的进媒体库**。

对应的问题：agent 无法和媒体库联动 —— 模型只能去 `execute_python` 里写宿主路径
（被沙箱正确拒绝），而 `media_create` 又只能传纯文本。修好之后这里钉住四件事：

1. **二进制可入媒体库**（word/Excel/图片走 base64），且 `type` 按扩展名推断 ——
   否则 .docx 会被标成 text，用户在「文档」筛选里找不到；
2. **落库失败必须显式说出来**（回退本地 store 时带 `stored_in` + `warning`）——
   此前是静默回退 + 报"创建成功"，这正是"看不见"的成因；
3. **落库列的硬长度**（name/tags）在本地就校验，不让 INSERT 失败退化成静默回退；
4. **体积上限**与网关内部端点一致，超限直接给出可操作的错误。
"""
from __future__ import annotations

import base64
import io
from pathlib import Path

import pytest

from app.media.store import LocalStore
from app.tools import media as media_mod
from app.tools.media import (
    MAX_ASSET_BYTES,
    MAX_NAME_CHARS,
    media_create,
)


@pytest.fixture()
def store(tmp_path, monkeypatch):
    """把引擎本地回退 store 指到临时目录（真实实现，不 mock）。"""
    local = LocalStore(str(tmp_path / "media"))
    monkeypatch.setattr(media_mod, "_store", local)
    return local


@pytest.fixture()
def library(monkeypatch):
    """假媒体库：记录调用参数，返回一个成功的资产。"""
    calls: list[dict] = []

    async def _persist(name, data, *, asset_type, category="generated", mime_type="", tags=None):
        calls.append({
            "name": name, "data": data, "asset_type": asset_type,
            "category": category, "mime_type": mime_type, "tags": list(tags or []),
        })
        return {
            "id": "asset-1",
            "name": name,
            "type": asset_type,
            "file_url": f"http://gateway/media/t1/asset-1/{name}",
            "size": len(data),
        }, ""

    monkeypatch.setattr(media_mod, "persist_to_library", _persist)
    return calls


@pytest.fixture()
def library_down(monkeypatch):
    """假媒体库：明确拒绝（模拟网关 400 —— 例如它拒了脚本类扩展名）。"""
    async def _persist(name, data, *, asset_type, category="generated", mime_type="", tags=None):
        return None, "media library rejected the asset (HTTP 400): file type not allowed: text/plain"

    monkeypatch.setattr(media_mod, "persist_to_library", _persist)


# ── 1. 文本资产：走媒体库，类型按扩展名推断 ──


@pytest.mark.asyncio
async def test_text_asset_lands_in_library(library, store):
    result = await media_create(name="summary.md", content="# 标题\n正文")

    assert result["stored_in"] == "library"
    assert result["type"] == "text"
    assert result["id"] == "asset-1"
    assert library[0]["asset_type"] == "text"
    assert library[0]["mime_type"] == "text/markdown"
    assert library[0]["data"] == "# 标题\n正文".encode()
    assert store.list() == []          # 没有落到本地回退路径


@pytest.mark.asyncio
async def test_csv_is_typed_as_text_not_default_default(library, store):
    result = await media_create(name="data.csv", content="a,b\n1,2\n")
    assert result["type"] == "text"
    assert library[0]["mime_type"] == "text/csv"


# ── 2. 二进制：这才是"创建一个 word 文件"的路 ──


@pytest.mark.asyncio
async def test_docx_bytes_reach_the_library_with_document_type(library, store):
    """用 python-docx 真造一份 .docx，走 base64 送进去。"""
    import docx

    document = docx.Document()
    document.add_heading("季度报告", level=1)
    document.add_paragraph("正文内容")
    buffer = io.BytesIO()
    document.save(buffer)
    raw = buffer.getvalue()

    result = await media_create(
        name="报告.docx",
        content_base64=base64.b64encode(raw).decode("ascii"),
    )

    assert result["stored_in"] == "library"
    # 关键：不是 "text" —— 否则它在「文档」筛选里永远找不到
    assert result["type"] == "document"
    assert library[0]["data"] == raw
    assert library[0]["mime_type"].endswith("wordprocessingml.document")
    assert result["size"] == len(raw)


@pytest.mark.asyncio
async def test_base64_tolerates_linebreaks_urlsafe_and_missing_padding(library, store):
    raw = bytes(range(256)) * 3
    encoded = base64.b64encode(raw).decode("ascii")
    # 模型生成的常见形态：分块换行 + URL-safe 字母表 + 去掉 padding
    messy = "\n".join(encoded[i:i + 60] for i in range(0, len(encoded), 60))
    messy = messy.replace("+", "-").replace("/", "_").rstrip("=")

    result = await media_create(name="blob.bin", content_base64=messy)

    assert result["stored_in"] == "library"
    assert library[0]["data"] == raw


# ── 3. 落库失败：必须显式 ──


@pytest.mark.asyncio
async def test_library_rejection_is_reported_not_hidden(library_down, store):
    result = await media_create(name="tool.py", content="print(1)")

    assert result["stored_in"] == "engine-local"
    assert "NOT in the media library" in result["output"]
    assert "file type not allowed" in result["warning"]   # 网关原文要带给模型
    assert result["file_url"]                       # 仍然可用（本地回退）
    assert len(store.list()) == 1


@pytest.mark.asyncio
async def test_success_never_mentions_engine_local(library, store):
    result = await media_create(name="ok.txt", content="hi")
    assert "stored_in" in result and result["stored_in"] == "library"
    assert "warning" not in result


# ── 4. 参数与体积校验（都在本地拦，不让 INSERT / 网关失败退化成静默回退）──


@pytest.mark.asyncio
async def test_oversized_content_is_rejected_before_calling_the_library(monkeypatch):
    called = False

    async def _persist(*_a, **_k):
        nonlocal called
        called = True
        return {}, ""

    monkeypatch.setattr(media_mod, "persist_to_library", _persist)
    result = await media_create(name="big.bin", content="x" * (MAX_ASSET_BYTES + 1))

    assert "too large" in result["error"]
    assert called is False


@pytest.mark.asyncio
async def test_name_validation(library, store):
    assert "required" in (await media_create(name="  ", content="x"))["error"]
    assert "too long" in (await media_create(name="a" * (MAX_NAME_CHARS + 1), content="x"))["error"]
    assert "path separators" in (await media_create(name="a/b.txt", content="x"))["error"]
    assert library == []


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "kwargs",
    [
        {"content": "a", "content_base64": "YQ=="},
        {"content": "a", "from_workspace": "x.txt"},
        {"content_base64": "YQ==", "from_workspace": "x.txt"},
        {"content": "a", "content_base64": "YQ==", "from_workspace": "x.txt"},
    ],
)
async def test_the_three_sources_are_mutually_exclusive(library, store, kwargs):
    """`content` / `content_base64` / `from_workspace` **只能给一个**。

    三种入参各有各的语义（文本 / 二进制 / 工作区文件），同时给两个时"以谁为准"没有合理答案
    —— 与其猜，不如让调用方说清楚（否则就是一次静默的数据来源覆盖）。
    """
    result = await media_create(name="x.txt", **kwargs)

    assert "only one of" in result["error"], result
    assert library == []


@pytest.mark.asyncio
async def test_unknown_type_is_rejected(library, store):
    result = await media_create(name="x.txt", content="a", type="spreadsheet")
    assert "unsupported type" in result["error"]


# ── 工作区 → 媒体库的发布通道（2026-10-09）──────────────────────────────────
#
# 这一路的动机：**先生成、再发布**。此前用 `execute_python` 在工作区生成 docx/xlsx 之后，
# 只能把字节编成 base64 再递回 `media_create`（大文件既慢又容易写坏），或者干脆写宿主路径
# （被沙箱正确拒绝）。现在直接指名工作区里的那个文件即可。


@pytest.fixture()
def workspace(tmp_path, monkeypatch):
    """把沙箱根指到 tmp：让"工作区"在用例里可读可写，且不污染宿主。"""
    from app.tools import sandbox

    monkeypatch.setenv(sandbox.SANDBOX_ROOT_ENV, str(tmp_path / "sandbox"))
    ws = sandbox.workspace_dir()
    ws.mkdir(parents=True, exist_ok=True)
    return ws


@pytest.mark.asyncio
async def test_publish_a_workspace_file(library, store, workspace):
    """`from_workspace` 把工作区里的文件**按原始字节**发布（不再要求 base64）。"""
    payload = b"PK\x03\x04 fake docx bytes"
    (workspace / "q3.docx").write_bytes(payload)

    result = await media_create(name="", from_workspace="q3.docx")

    assert "error" not in result, result
    assert result["name"] == "q3.docx", "名字可以从工作区路径兜底"
    assert result["size"] == len(payload)
    assert library and library[0]["data"] == payload


@pytest.mark.asyncio
async def test_publish_keeps_the_workspace_file(library, store, workspace):
    """发布是**读**，不是搬移：工作区里的文件还在（后续步骤可能继续引用它）。"""
    (workspace / "keep.md").write_text("hello", encoding="utf-8")

    await media_create(name="keep.md", from_workspace="keep.md")

    assert (workspace / "keep.md").exists()


@pytest.mark.asyncio
async def test_workspace_path_cannot_escape(library, store, workspace):
    """路径必须 clamp 在工作区内 —— 否则这个参数就是一条**任意文件读**通道。

    `../` 逃逸在 Windows 与 POSIX 上语义一致，因此拿它钉"**被拒绝**"（而不是"碰巧不存在"；
    只断言"报错"是弱判据：文件不存在同样报错）。
    """
    result = await media_create(name="x.txt", from_workspace="../secret.txt")
    assert "outside the sandbox workspace" in result["error"], result

    for bad in ("..\\secret.txt", "/etc/passwd", "C:\\Windows\\win.ini"):
        result = await media_create(name="x.txt", from_workspace=bad)
        assert "error" in result, bad
    assert library == []


@pytest.mark.asyncio
async def test_missing_workspace_file_is_explicit(library, store, workspace):
    result = await media_create(name="x.txt", from_workspace="nope.txt")

    assert "no such file in the workspace" in result["error"]
    assert "generate it first" in result["error"], "错误里要给出**下一步**怎么办"


@pytest.mark.asyncio
async def test_directory_is_refused(library, store, workspace):
    (workspace / "sub").mkdir()

    result = await media_create(name="x.txt", from_workspace="sub")

    assert "is a directory" in result["error"]


@pytest.mark.asyncio
async def test_size_guard_runs_before_reading(library, store, workspace, monkeypatch):
    """体积上限必须在**消费之前**生效：超限时连 `read_bytes` 都不该被调用。

    只断言"报错"不够 —— 读完再判大小同样会报错，但那时整个文件已经进了内存。
    （§4 的同族教训：「有上限」不等于「在消费前拦住」。）
    """
    monkeypatch.setattr(media_mod, "MAX_ASSET_BYTES", 16)
    (workspace / "big.bin").write_bytes(b"x" * 64)

    reads: list[str] = []
    original = Path.read_bytes

    def _spy(self):
        reads.append(str(self))
        return original(self)

    monkeypatch.setattr(Path, "read_bytes", _spy)

    result = await media_create(name="big.bin", from_workspace="big.bin")

    assert "too large" in result["error"]
    assert reads == [], "超限时不得把文件读进内存"


@pytest.mark.asyncio
async def test_tags_are_capped_to_column_width(library, store):
    result = await media_create(
        name="x.md",
        content="a",
        tags=["t" * 80, "dup", "dup"] + [f"tag{i}" for i in range(50)],
    )
    assert result["stored_in"] == "library"
    sent = library[0]["tags"]
    assert len(",".join(sent)) <= media_mod.MAX_TAGS_TOTAL_CHARS
    assert all(len(t) <= media_mod.MAX_TAG_CHARS for t in sent)
    assert len(sent) == len(set(sent))              # 去重


# ── 5. 纯函数 ──


@pytest.mark.parametrize("mime,name,expected", [
    ("image/png", "a.png", "image"),
    ("audio/mpeg", "a.mp3", "audio"),
    ("video/mp4", "a.mp4", "video"),
    ("", "报告.docx", "document"),
    ("", "book.pdf", "document"),
    ("application/zip", "data.xlsx", "document"),
    ("text/plain", "note.md", "text"),
    ("application/json", "payload.bin", "text"),
    ("application/octet-stream", "blob.bin", "file"),
    ("", "script.unknown", "text"),
])
def test_infer_asset_type(mime, name, expected):
    assert media_mod._infer_asset_type(mime, name) == expected


def test_registered_schema_exposes_binary_path():
    """工具 schema 必须真的把二进制入口暴露给模型（否则它只能去写宿主路径）。"""
    from app.tools.registry import registry

    definition = registry.get("media_create")
    assert definition is not None
    properties = definition.parameters["properties"]
    assert "content_base64" in properties
    assert "mime_type" in properties
    assert definition.parameters["required"] == ["name"]
    assert "media library" in definition.description


def test_media_create_is_visible_to_the_model_by_default():
    """核心工具集必须含 media_create —— 注册了却默认不可见，模型必然走错路。"""
    from app.agent.modes import CORE_TOOL_NAMES

    assert "media_create" in CORE_TOOL_NAMES
